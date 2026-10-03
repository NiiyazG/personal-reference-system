"""Coverage for stage-1 code paths that shipped without a test.

Each test here exercises behaviour that was already implemented but never
verified, so a silent regression in these branches would go unnoticed.
"""

import sys
import tempfile
import unittest
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from reference_system.core import ReferenceService
from reference_system.foundation import initialize_project


class _TempProject:
    def __enter__(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name) / "personal-reference"
        initialize_project(self.root)
        self.workspace = Path(self._tmp.name)
        return self

    def __exit__(self, *exc):
        self._tmp.cleanup()
        return False


class EncodingFallbackTests(unittest.TestCase):
    def test_cp1251_text_is_decoded_and_searchable(self):
        with _TempProject() as project:
            source = project.workspace / "legacy.txt"
            source.write_bytes("Насос вибрирует при пуске.\n".encode("cp1251"))

            service = ReferenceService(project.root)
            added = service.add_file(source)

            self.assertEqual(added["detected_type"], "txt")
            hits = service.search("вибрирует")
            self.assertEqual(len(hits), 1)
            self.assertIn("вибрирует", hits[0]["text"])

    def test_cp1251_csv_is_decoded_and_searchable(self):
        with _TempProject() as project:
            source = project.workspace / "legacy.csv"
            source.write_bytes("id;описание\n1;Задвижка DN100\n".encode("cp1251"))

            service = ReferenceService(project.root)
            added = service.add_file(source)

            self.assertEqual(added["detected_type"], "csv")
            hits = service.search("Задвижка", exact=True)
            self.assertEqual(len(hits), 1)
            self.assertEqual(hits[0]["locator"]["type"], "csv-cell")
            self.assertEqual(hits[0]["locator"]["row"], 2)
            self.assertEqual(hits[0]["locator"]["column"], 2)

    def test_detected_encoding_is_reported_honestly(self):
        from reference_system.processors import extract_file

        with _TempProject() as project:
            legacy = project.workspace / "legacy.txt"
            legacy.write_bytes("Насос вибрирует.\n".encode("cp1251"))
            self.assertEqual(extract_file(legacy)["metadata"]["encoding"], "cp1251")

            modern = project.workspace / "modern.txt"
            modern.write_text("Насос вибрирует.\n", encoding="utf-8")
            self.assertEqual(extract_file(modern)["metadata"]["encoding"], "utf-8")

            dos = project.workspace / "dos.txt"
            dos.write_bytes("Насос вибрирует.\n".encode("cp866"))
            self.assertIn(extract_file(dos)["metadata"]["encoding"], {"cp866", "cp1251"})

    def test_binary_text_file_is_rejected(self):
        with _TempProject() as project:
            source = project.workspace / "blob.txt"
            source.write_bytes(b"\x00\x01\x02\x00binary")
            service = ReferenceService(project.root)
            with self.assertRaises(Exception) as ctx:
                service.add_file(source)
            self.assertIn("NUL", str(ctx.exception))

    def test_utf8_bom_and_crlf_are_normalised(self):
        with _TempProject() as project:
            source = project.workspace / "windows.txt"
            source.write_bytes("\ufeffПервая строка\r\nВторая строка\r\n".encode("utf-8"))

            service = ReferenceService(project.root)
            service.add_file(source)
            hit = service.search("Первая")[0]

            self.assertNotIn("\ufeff", hit["text"])
            self.assertNotIn("\r", hit["text"])
            self.assertEqual(hit["locator"]["line_end"], 2)
            self.assertEqual(hit["locator"]["line_start"], 1)


class CsvDialectTests(unittest.TestCase):
    def _add(self, project: "_TempProject", name: str, payload: str) -> ReferenceService:
        source = project.workspace / name
        source.write_text(payload, encoding="utf-8")
        service = ReferenceService(project.root)
        service.add_file(source)
        return service

    def test_pipe_separated_file_is_split_on_pipes(self):
        with _TempProject() as project:
            service = self._add(project, "data.csv", "id|описание\n1|Клапан обратный\n")
            hit = service.search("Клапан", exact=True)[0]
            self.assertEqual(hit["locator"]["column"], 2)

    def test_tsv_extension_is_accepted_and_split_on_tabs(self):
        with _TempProject() as project:
            service = self._add(project, "data.tsv", "id\tописание\n1\tНасос P-101\n")
            hit = service.search("Насос", exact=True)[0]
            self.assertEqual(hit["locator"]["type"], "csv-cell")
            self.assertEqual(hit["locator"]["column"], 2)
            self.assertEqual(hit["locator"]["row"], 2)

    def test_tab_extension_is_accepted(self):
        with _TempProject() as project:
            service = self._add(project, "data.tab", "id\tописание\n1\tМуфта упругая\n")
            self.assertEqual(service.search("Муфта", exact=True)[0]["locator"]["column"], 2)

    def test_comma_and_semicolon_are_both_detected(self):
        with _TempProject() as project:
            service = self._add(project, "a.csv", "id,описание\n1,Муфта упругая\n")
            self.assertEqual(service.search("Муфта", exact=True)[0]["locator"]["column"], 2)

        with _TempProject() as project:
            service = self._add(project, "b.csv", "id;описание\n1;Муфта упругая\n")
            self.assertEqual(service.search("Муфта", exact=True)[0]["locator"]["column"], 2)


class OfficeBranchTests(unittest.TestCase):
    def test_xlsx_formula_cell_is_indexed_as_the_formula(self):
        import openpyxl

        with _TempProject() as project:
            source = project.workspace / "calc.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet["A1"] = 2
            sheet["A2"] = 3
            sheet["B1"] = "=SUM(A1:A2)"
            workbook.save(source)
            workbook.close()

            service = ReferenceService(project.root)
            added = service.add_file(source)
            self.assertEqual(added["detected_type"], "xlsx")

            hits = service.search("SUM", exact=True)
            self.assertEqual(len(hits), 1)
            self.assertEqual(hits[0]["locator"]["cell"], "B1")
            self.assertEqual(hits[0]["locator"]["sheet"], "Sheet")

    def test_docx_cell_with_several_paragraphs_is_one_fragment(self):
        import docx

        with _TempProject() as project:
            source = project.workspace / "table.docx"
            document = docx.Document()
            table = document.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            cell.text = "Первое требование."
            cell.add_paragraph("Второе требование.")
            document.save(source)

            service = ReferenceService(project.root)
            service.add_file(source)

            hit = service.search("Второе", exact=True)[0]
            self.assertEqual(hit["locator"]["type"], "docx-table-cell")
            self.assertEqual(hit["locator"]["table"], 1)
            self.assertEqual(hit["locator"]["row"], 1)
            self.assertEqual(hit["locator"]["column"], 1)
            self.assertIn("Первое требование.", hit["text"])


class PdfPageTests(unittest.TestCase):
    def test_pdf_with_one_text_page_and_one_blank_page_indexes_only_the_text(self):
        import fitz

        with _TempProject() as project:
            source = project.workspace / "mixed.pdf"
            document = fitz.open()
            first = document.new_page()
            first.insert_text((72, 72), "Pump P-101 inspection record")
            document.new_page()
            document.save(source)
            document.close()

            service = ReferenceService(project.root)
            added = service.add_file(source)

            self.assertEqual(added["detected_type"], "pdf")
            hits = service.search("inspection")
            self.assertEqual(len(hits), 1)
            self.assertEqual(hits[0]["locator"]["type"], "pdf-page")
            self.assertEqual(hits[0]["locator"]["page"], 1)

    def test_pdf_with_only_blank_pages_is_rejected(self):
        import fitz

        with _TempProject() as project:
            source = project.workspace / "blank.pdf"
            document = fitz.open()
            document.new_page()
            document.new_page()
            document.save(source)
            document.close()

            service = ReferenceService(project.root)
            with self.assertRaises(Exception) as ctx:
                service.add_file(source)
            self.assertIn("OCR", str(ctx.exception))


class ScaleAndLimitTests(unittest.TestCase):
    def test_limit_is_clamped_and_integrity_holds_at_scale(self):
        with _TempProject() as project:
            service = ReferenceService(project.root)
            for index in range(120):
                source = project.workspace / f"item-{index:03d}.txt"
                source.write_text(f"Запись номер {index:03d} содержит маркер SEEKME.\n", encoding="utf-8")
                service.add_file(source)

            self.assertEqual(len(service.search("SEEKME", exact=True, limit=500)), 100)
            self.assertEqual(len(service.search("SEEKME", exact=True, limit=5)), 5)
            self.assertEqual(len(service.search("SEEKME", exact=True, limit=0)), 1)

            status = service.status()
            self.assertEqual(status["materials"], 120)
            self.assertEqual(status["fragments"], 120)

            report = service.integrity_check()
            self.assertTrue(report["ok"], report["issues"])

    def test_reprocess_after_reprocess_keeps_a_single_fragment_set(self):
        with _TempProject() as project:
            source = project.workspace / "note.txt"
            source.write_text("Редуктор требует замены масла.\n", encoding="utf-8")

            service = ReferenceService(project.root)
            added = service.add_file(source)
            service.reprocess(added["material_id"])
            service.reprocess(added["material_id"])

            self.assertEqual(len(service.search("Редуктор")), 1)
            self.assertEqual(service.status()["fragments"], 1)
            self.assertTrue(service.integrity_check()["ok"])


if __name__ == "__main__":
    unittest.main()
