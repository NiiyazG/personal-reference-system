import sys
import tempfile
import unittest
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from reference_system.core import ReferenceService
from reference_system.foundation import initialize_project
from reference_system.processors import UnsupportedFormatError


class ProcessorTests(unittest.TestCase):
    def test_pdf_text_layer_is_searchable_with_page_locator(self):
        import fitz

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "personal-reference"
            initialize_project(root)
            source = Path(tmp) / "manual.pdf"
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((72, 72), "Pump service interval 500 hours")
            doc.save(source)
            doc.close()

            service = ReferenceService(root)
            added = service.add_file(source, title="Pump Manual")

            self.assertEqual(added["detected_type"], "pdf")
            hits = service.search("service interval")
            self.assertEqual(len(hits), 1)
            source_ref = service.show_source(hits[0]["fragment_id"])
            self.assertEqual(source_ref["locator"]["type"], "pdf-page")
            self.assertEqual(source_ref["locator"]["page"], 1)

    def test_docx_table_text_is_searchable_with_cell_locator(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "personal-reference"
            initialize_project(root)
            source = Path(tmp) / "specification.docx"
            doc = Document()
            doc.add_paragraph("Equipment specification")
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Parameter"
            table.cell(0, 1).text = "Nominal flow 42 m3/h"
            doc.save(source)

            service = ReferenceService(root)
            added = service.add_file(source)

            self.assertEqual(added["detected_type"], "docx")
            hits = service.search("nominal flow")
            self.assertEqual(len(hits), 1)
            source_ref = service.show_source(hits[0]["fragment_id"])
            self.assertEqual(source_ref["locator"]["type"], "docx-table-cell")
            self.assertEqual(source_ref["locator"]["table"], 1)
            self.assertEqual(source_ref["locator"]["row"], 1)
            self.assertEqual(source_ref["locator"]["column"], 2)

    def test_xlsx_cell_is_searchable_with_sheet_and_cell_locator(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "personal-reference"
            initialize_project(root)
            source = Path(tmp) / "measurements.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Data"
            sheet["A1"] = "Metric"
            sheet["B2"] = "Outlet temperature 68 C"
            workbook.save(source)
            workbook.close()

            service = ReferenceService(root)
            added = service.add_file(source)

            self.assertEqual(added["detected_type"], "xlsx")
            hits = service.search("outlet temperature")
            self.assertEqual(len(hits), 1)
            source_ref = service.show_source(hits[0]["fragment_id"])
            self.assertEqual(source_ref["locator"]["type"], "xlsx-cell")
            self.assertEqual(source_ref["locator"]["sheet"], "Data")
            self.assertEqual(source_ref["locator"]["cell"], "B2")

    def test_csv_cell_is_searchable_with_row_and_column_locator(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "personal-reference"
            initialize_project(root)
            source = Path(tmp) / "assets.csv"
            source.write_text("id;description\n1;Heat exchanger plate A17\n", encoding="utf-8")

            service = ReferenceService(root)
            added = service.add_file(source)

            self.assertEqual(added["detected_type"], "csv")
            hits = service.search("heat exchanger")
            self.assertEqual(len(hits), 1)
            source_ref = service.show_source(hits[0]["fragment_id"])
            self.assertEqual(source_ref["locator"]["type"], "csv-cell")
            self.assertEqual(source_ref["locator"]["row"], 2)
            self.assertEqual(source_ref["locator"]["column"], 2)

    def test_png_is_registered_with_metadata_without_ocr(self):
        from PIL import Image

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "personal-reference"
            initialize_project(root)
            source = Path(tmp) / "inspection.png"
            Image.new("RGB", (32, 18), color=(210, 220, 200)).save(source)

            service = ReferenceService(root)
            added = service.add_file(source, title="Inspection Image")

            self.assertEqual(added["detected_type"], "png")
            hits = service.search("Inspection Image")
            self.assertEqual(len(hits), 1)
            source_ref = service.show_source(hits[0]["fragment_id"])
            self.assertEqual(source_ref["locator"]["type"], "image-file")
            self.assertEqual(source_ref["locator"]["width"], 32)
            self.assertEqual(source_ref["locator"]["height"], 18)
            self.assertFalse(source_ref["locator"]["ocr_performed"])

    def test_jpeg_is_registered_with_dimensions_without_ocr(self):
        from PIL import Image

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "personal-reference"
            initialize_project(root)
            source = Path(tmp) / "pump.jpg"
            Image.new("RGB", (40, 24), color=(190, 195, 205)).save(source, format="JPEG")

            service = ReferenceService(root)
            added = service.add_file(source, title="Pump Photo")

            self.assertEqual(added["detected_type"], "jpeg")
            hits = service.search("Pump Photo")
            self.assertEqual(len(hits), 1)
            source_ref = service.show_source(hits[0]["fragment_id"])
            self.assertEqual(source_ref["locator"]["type"], "image-file")
            self.assertEqual(source_ref["locator"]["width"], 40)
            self.assertEqual(source_ref["locator"]["height"], 24)
            self.assertEqual(source_ref["locator"]["format"], "jpeg")
            self.assertFalse(source_ref["locator"]["ocr_performed"])

    def test_image_with_wrong_signature_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "personal-reference"
            initialize_project(root)
            source = Path(tmp) / "fake.png"
            source.write_bytes(b"not really a png at all")

            service = ReferenceService(root)
            with self.assertRaises(UnsupportedFormatError):
                service.add_file(source)

    def test_processors_depend_on_content_not_on_file_name(self):
        import shutil

        import docx
        import openpyxl

        from reference_system.processors import extract_file

        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp)
            docx_source = workspace / "manual.docx"
            document = docx.Document()
            document.add_paragraph("Порядок пуска насоса.")
            document.save(docx_source)

            xlsx_source = workspace / "log.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active["A1"] = "Температура"
            workbook.save(xlsx_source)
            workbook.close()

            for source, original_name, expected_type in (
                (docx_source, "manual.docx", "docx"),
                (xlsx_source, "log.xlsx", "xlsx"),
            ):
                anonymous = workspace / "input.bin"
                shutil.copyfile(source, anonymous)
                extracted = extract_file(anonymous, original_name=original_name)
                self.assertEqual(extracted["detected_type"], expected_type)
                self.assertTrue(extracted["fragments"])
                anonymous.unlink()


if __name__ == "__main__":
    unittest.main()
